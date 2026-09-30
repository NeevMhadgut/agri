import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from writeup_helpers import (
    add_header_banner, add_meta_table, add_boxed_section,
    add_code_block, add_screenshot
)

HEADER_IMG = "images/somaiya_header.png"
os.makedirs("writeups", exist_ok=True)

def generate_exp1():
    doc = Document()
    add_header_banner(doc, HEADER_IMG)
    add_meta_table(doc, exp_no="1",
                   title="Design and Development of an AI-based Smart Irrigation Advisory Web Portal Front-End using HTML5",
                   date_perf="15 / 07 / 2026")

    # Aim
    add_boxed_section(doc, "Aim of the Experiment:",
        "To design and develop a multi-page, use-case-based web portal front-end using HTML5 for an AI-based "
        "Smart Irrigation Advisory System for sugarcane crops (KJS-AGR-01), incorporating semantic structure, "
        "lists, text formatting, multimedia, interactive client-side image maps, responsive tables, and forms.")

    # Objectives
    add_boxed_section(doc, "Objectives of the Experiment:", [
        "1. To develop structured and semantic HTML5 web pages for an AI-enabled smart irrigation advisory system.",
        "2. To create interactive web pages for farmer registration, irrigation recommendations, and weather forecast display.",
        "3. To apply real-world agricultural IoT use case telemetry as the core theme across all modular pages.",
        "4. To implement HTML5 lists, tables, image maps, forms, and multimedia elements without external libraries.",
        "5. To build a robust, accessible foundation ready for subsequent CSS3 styling and JavaScript validation layers."
    ])

    # COs
    add_boxed_section(doc, "COs to be achieved:", "CO1: Developing webpages using HTML and CSS.")

    # References
    add_boxed_section(doc, "Books / Journals / Websites references:", [
        "1. MDN Web Docs – HTML: HyperText Markup Language, Mozilla Developer Network, 2026. https://developer.mozilla.org/en-US/docs/Web/HTML",
        "2. W3Schools – HTML5 Tutorial and Semantic Web Guide, Refsnes Data, 2026. https://www.w3schools.com/html/",
        "3. J. Duckett, HTML and CSS: Design and Build Websites, 1st ed., Indianapolis, IN: John Wiley & Sons, 2011."
    ])

    # Theory
    theory_text = (
        "HyperText Markup Language 5 (HTML5) is the core standard for structuring and presenting content on the World Wide Web [1]. "
        "HTML5 introduces semantic tags that provide explicit meaning to both browsers and search engines, superseding legacy generic '<div>' elements.\n\n"
        "Key HTML5 Modules Used in This Portal:\n"
        "1. Semantic Structural Elements: Elements including '<header>', '<nav>', '<main>', '<section>', '<article>', '<aside>', and '<footer>' "
        "delineate the layout hierarchy and ensure accessibility.\n"
        "2. Information Representation via Lists: Ordered lists ('<ol>'), unordered lists ('<ul>'), nested lists, and description lists ('<dl>') "
        "organize sensor telemetry collection sequences, agricultural parameters, and domain abbreviations.\n"
        "3. Tabular Data Structures: Complex tables utilizing '<thead>', '<tbody>', '<tfoot>', '<th>', '<caption>', along with 'rowspan' and 'colspan' "
        "attributes present multi-zone sugarcane soil moisture thresholds, irrigation schedules, and weather dependencies.\n"
        "4. Client-Side Image Maps: The '<map>' and '<area>' tags enable coordinates-based interactive navigation on aerial farm plot imagery.\n"
        "5. Comprehensive HTML5 Forms: Semantic input controls ('text', 'tel', 'email', 'number', 'date', 'select', 'textarea') enclosed within "
        "'<fieldset>' and '<legend>' containers enable structured data acquisition.\n"
        "6. Embedded Media & Iframes: Native '<audio>' and '<video>' elements deliver farmer training media without third-party plugins, "
        "while '<iframe>' embeds georeferenced GIS telemetry radar."
    )
    add_boxed_section(doc, "Theory / Relevant Concepts:", theory_text)

    # Step by Step Procedure
    proc_text = (
        "1. Define the website wireframe and modular multi-page hierarchy for the Sugarcane Advisory System (KJS-AGR-01).\n"
        "2. Task 1: Construct the primary portal layout in 'index.html' using semantic tags (<header>, <nav>, <main>, <aside>, <footer>).\n"
        "3. Task 2: Create 'farm-information.html' to present telemetry workflows and abbreviations using ordered, unordered, and definition lists.\n"
        "4. Task 3: Build 'irrigation-advisory.html' employing HTML formatting tags (<strong>, <em>, <mark>, <abbr>, <sub>, <sup>).\n"
        "5. Task 4: Construct the precision agriculture gallery in 'index.html' using semantic <figure>, <img>, and <figcaption> tags.\n"
        "6. Task 5: Implement 'farm-map.html' with a client-side <map> and rectangular/polygonal <area> coordinates linking to plot advisories.\n"
        "7. Task 6: Build 'recommendation.html' with structured <table>, <thead>, <tbody>, <tfoot>, rowspan, and colspan attributes.\n"
        "8. Task 7: Design the farmer registration form in 'registration.html' with required fieldsets, inputs, and validation attributes.\n"
        "9. Task 8: Embed an OpenStreetMap GIS radar iframe into 'media.html' for live field coordinates tracking.\n"
        "10. Task 9: Embed awareness video and audio broadcasting elements in 'media.html' using native HTML5 media tags."
    )
    add_boxed_section(doc, "Step by Step Implementation / Procedure:", proc_text)

    # Code Listings
    doc.add_page_break()
    p_code = doc.add_paragraph()
    r = p_code.add_run("HTML5 Source Code Snippets:")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(128, 0, 0)

    # Task 1
    add_code_block(doc, "Task 1: Semantic Homepage Structure (index.html)", """<!-- index.html excerpt -->
<header class="layout-header">
  <p class="org-name"><img class="logo" src="./images/image.png" alt="org logo"> Agricultural Advisory Organization</p>
  <h1 class="page-name">Smart Irrigation Advisory System</h1>
</header>
<nav class="site-nav">
  <ul class="nav-menu">
    <li><a href="index.html">Home</a></li>
    <li><a href="irrigation-advisory.html">Irrigation Advisory</a></li>
    <li><a href="farm-information.html">Farm Information</a></li>
    <li><a href="farm-map.html">Interactive Map</a></li>
    <li><a href="recommendation.html">Recommendations</a></li>
    <li><a href="registration.html">Farmer Registration</a></li>
    <li><a href="media.html">Media & Awareness</a></li>
  </ul>
</nav>""")

    # Task 2
    add_code_block(doc, "Task 2: Agricultural Lists Implementation (farm-information.html)", """<!-- farm-information.html excerpt -->
<ol>
  <li>Collect farm profile and sugarcane crop stage details.</li>
  <li>Read telemetry from IoT soil moisture sensors.</li>
  <li>Analyze soil moisture deficit against crop evapotranspiration coefficients.</li>
  <li>Generate plot-specific irrigation dosage and run-time recommendations.</li>
</ol>
<dl>
  <dt>AI</dt><dd>Artificial Intelligence — Computational models simulating agronomic decisions.</dd>
  <dt>IoT</dt><dd>Internet of Things — Network of physical FDR probes collecting moisture telemetry.</dd>
</dl>""")

    # Task 5 & 6
    add_code_block(doc, "Task 5 & 6: Image Map and Tabular Data Structures", """<!-- Farm Map Image Map (farm-map.html) -->
<img src="./images/farm_map.jpg" alt="Sugarcane farm map" usemap="#farmMap">
<map name="farmMap">
  <area shape="rect" coords="20,20,280,180" href="irrigation-advisory.html#plot-a" alt="Plot A note">
  <area shape="rect" coords="300,20,580,180" href="irrigation-advisory.html#plot-b" alt="Plot B note">
  <area shape="rect" coords="20,200,580,380" href="irrigation-advisory.html#plot-c" alt="Plot C note">
</map>

<!-- Table with rowspan and colspan (recommendation.html) -->
<table class="data-table">
  <thead>
    <tr><th>Plot ID</th><th>Farmer Name</th><th>Stage</th><th>Moisture</th><th>Duration</th></tr>
  </thead>
  <tbody>
    <tr><td rowspan="2">Plot A</td><td>Ravi Patil</td><td>Tillering</td><td>28%</td><td>35 min</td></tr>
    <tr><td>Ravi Patil</td><td colspan="3">Young crop requiring low moisture threshold</td></tr>
  </tbody>
</table>""")

    # Screenshots
    doc.add_page_break()
    p_out = doc.add_paragraph()
    r = p_out.add_run("Expected Output / Screenshots (Task-Wise):")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(128, 0, 0)

    add_screenshot(doc, "Task 1: Semantic Homepage Portal Structure (index.html)", "screenshots/exp1_task1_homepage.png")
    add_screenshot(doc, "Task 2: Farm Information Displayed using Ordered, Unordered & Definition Lists", "screenshots/exp1_task2_lists.png")
    add_screenshot(doc, "Task 3: Irrigation Advisory Page using HTML Formatting & Abbreviation Tags", "screenshots/exp1_task3_formatting.png")
    add_screenshot(doc, "Task 4: Farm Gallery using Semantic Figure, Image & Figcaption Elements", "screenshots/exp1_task4_gallery.png")
    add_screenshot(doc, "Task 5: Interactive Farm Plot Navigation using Client-Side Image Map (<map> & <area>)", "screenshots/exp1_task5_imagemap.png")
    add_screenshot(doc, "Task 6: Soil Moisture & Irrigation Quota Table (Thead, Tbody, Rowspan & Colspan)", "screenshots/exp1_task6_table.png")
    add_screenshot(doc, "Task 7: Farmer Registration Form with Semantic Controls & Fieldsets", "screenshots/exp1_task7_form.png")
    add_screenshot(doc, "Task 8: Display Weather Forecast & Georeferenced Radar using IFrame", "screenshots/exp1_task8_iframe.png")
    add_screenshot(doc, "Task 9: Multimedia & Executable Awareness Content (Native HTML5 Audio & Video)", "screenshots/exp1_task9_media.png")

    # Post Lab Questions
    post_lab_q = (
        "Question 1: Explain the importance of semantic elements in HTML5 compared to traditional div tags.\n\n"
        "Answer:\n"
        "Semantic HTML tags (such as <header>, <nav>, <article>, <section>, and <footer>) describe their meaning to both the browser and developer. "
        "Traditional <div> tags are generic block containers that convey no structural context. The advantages of semantic elements include:\n"
        "1. Accessibility (a11y): Screen readers navigate landmarks directly, allowing visually impaired users to jump across navigation and main content.\n"
        "2. Search Engine Optimization (SEO): Search web crawlers accurately index content priority, distinguishing headlines from auxiliary sidebars.\n"
        "3. Maintainability: Codebases are self-documenting and easier for distributed engineering teams to read and debug.\n\n"
        "Question 2: How does a client-side image map function?\n\n"
        "Answer:\n"
        "An HTML client-side image map utilizes the '<img>' tag paired with 'usemap=\"#mapName\"'. The companion '<map name=\"mapName\">' element "
        "contains one or more '<area>' tags defining geometric shapes ('rect', 'circle', 'poly') and pixel coordinates ('coords=\"x1,y1,x2,y2\"'). "
        "When a user clicks within these geometric bounds, the browser routes them to the associated hyperlink without requiring server coordinate processing."
    )
    add_boxed_section(doc, "Post Lab Subjective/Objective type Questions:", post_lab_q, bg_hex="F4F6F9")

    # Conclusion & Discussion
    conc_text = (
        "Conclusion:\n"
        "Experiment 1 successfully established the front-end foundation for the AI-based Smart Irrigation Advisory System (KJS-AGR-01) using pure HTML5. "
        "A multi-page website was implemented incorporating semantic page structuring, comprehensive lists, formatted agronomic text directives, "
        "client-side clickable image maps for farm plots, structured tables with cell spans, interactive forms, and native multimedia awareness components.\n\n"
        "Discussion:\n"
        "Developing the portal emphasized strict structural hierarchy before applying CSS styles. Semantic elements made document navigation clean "
        "and logical. The client-side image map proved highly practical for rural agricultural use, enabling farmers to tap directly on visual plot "
        "zones to retrieve plot-specific irrigation advice."
    )
    add_boxed_section(doc, "Conclusion and Discussion:", conc_text)

    doc.save("writeups/Experiment_1_HTML5_Writeup.docx")
    print("Saved writeups/Experiment_1_HTML5_Writeup.docx")

if __name__ == "__main__":
    generate_exp1()
